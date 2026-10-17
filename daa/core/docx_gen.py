"""Word 文档生成逻辑 — 将 Markdown 报告转换为 .docx。

报告在提示词层面被要求写成段落体（daa/core/prompts.py），因此转换以
「标题 + 段落 + 图片」为主，同时兼容列表 / 表格 / 代码围栏。

注意：python-docx 的 ``run.font.name`` 只写 ``w:ascii`` / ``w:hAnsi``（拉丁
字体），中文需要额外写 ``w:eastAsia``，否则中文字符会回退到默认字体渲染。
"""

import os
import re
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT_NAME = "微软雅黑"
MONO_FONT = "Consolas"
BODY_SIZE = 11
CODE_SIZE = 9
TABLE_SIZE = 9


def _set_run_font(
    run,
    name: str = FONT_NAME,
    size: Optional[int] = None,
    bold: Optional[bool] = None,
    mono: bool = False,
) -> None:
    """统一设置 run 的字体、字号、加粗。

    关键：``run.font.name`` 只写拉丁字体（w:ascii / w:hAnsi），中文需额外写
    ``w:eastAsia`` 的 rFonts，否则中文字符不会用指定字体渲染。
    """
    font = run.font
    font.name = MONO_FONT if mono else name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold

    try:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


def _set_style_font(style, name: str = FONT_NAME) -> None:
    """设置段落样式（Normal 等）的中文字体，让正文默认继承。"""
    try:
        style.font.name = name
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), name)
    except Exception:
        pass


_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def _add_inline_runs(paragraph, text: str, size: Optional[int] = None) -> None:
    """按 ``**加粗**`` 与 `` `代码` `` 内联标记拆分为多个 run 写入段落。"""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _set_run_font(paragraph.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            _set_run_font(paragraph.add_run(part[1:-1]), size=size, mono=True)
        else:
            _set_run_font(paragraph.add_run(part), size=size)


def _is_table_separator(line: str) -> bool:
    """判断是否为 Markdown 表格的分隔行（|---|:--:|）。"""
    return bool(re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", line))


def _parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _fill_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = text
    for run in cell.paragraphs[0].runs:
        _set_run_font(run, size=TABLE_SIZE, bold=bold)


def render_docx(markdown: str, figures_dir: str, output_path: str) -> str:
    """把 Markdown 渲染成 .docx（返回输出路径）。

    图片引用 ``![alt](./chart.png)`` 按 ``os.path.basename`` 在 figures_dir 下
    解析（与 figures 端点一致），文件不存在时静默跳过，不中断转换。
    """
    doc = Document()
    _set_style_font(doc.styles["Normal"])

    lines = markdown.splitlines()
    n = len(lines)
    i = 0
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码围栏
        if stripped.startswith("```"):
            code_lines: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # 跳过闭合围栏
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                _set_run_font(p.add_run(cl), size=CODE_SIZE, mono=True)
            continue

        # 图片
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", stripped)
        if m:
            img_path = os.path.join(figures_dir, os.path.basename(m.group(1).strip()))
            if os.path.isfile(img_path):
                try:
                    doc.add_picture(img_path, width=Inches(6))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    print(f"⚠️ Word 插入图片失败（{img_path}）: {e}")
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and _is_table_separator(lines[i + 1].strip()):
            header = _parse_table_row(stripped)
            rows: list[list[str]] = []
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(_parse_table_row(lines[i].strip()))
                i += 1
            if header:
                ncols = len(header)
                table = doc.add_table(rows=1 + len(rows), cols=ncols)
                table.style = "Table Grid"
                for c, cell_text in enumerate(header):
                    _fill_cell(table.rows[0].cells[c], cell_text, bold=True)
                for r, row in enumerate(rows, start=1):
                    for c in range(ncols):
                        _fill_cell(table.rows[r].cells[c], row[c] if c < len(row) else "")
            continue

        # 标题（>3 级一律收敛到 Heading 3）
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            heading = doc.add_heading(level=min(len(m.group(1)), 3))
            _set_run_font(heading.add_run(m.group(2)))
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*•]\s+(.*)$", stripped)
        if m:
            _add_inline_runs(doc.add_paragraph(style="List Bullet"), m.group(1), size=BODY_SIZE)
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            _add_inline_runs(doc.add_paragraph(style="List Number"), m.group(1), size=BODY_SIZE)
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        _add_inline_runs(doc.add_paragraph(), stripped, size=BODY_SIZE)
        i += 1

    doc.save(output_path)
    return output_path
