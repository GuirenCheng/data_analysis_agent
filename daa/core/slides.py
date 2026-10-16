"""PPT 幻灯片生成逻辑 — LLM 提炼大纲 + python-pptx 渲染。

设计原因：最终分析报告被提示词刻意要求写成段落体散文（避免分点列表），
直接按 Markdown 结构映射成幻灯片会变成文字密集页。因此这里让 LLM 把报告
浓缩成「标题页 + 分节要点 + 配图页」的结构化大纲，再由 python-pptx 渲染成
真正的 .pptx，并嵌入分析过程已生成的图表图片。

大纲四类 slide：
    - title      : title + subtitle
    - section    : title + bullets[]
    - figure     : title + figure(精确文件名) + caption
    - conclusion : title + bullets[]
"""

import os
from typing import Any, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from daa.core.protocol import parse_yaml_response
from daa.core.reporter import build_figures_summary

# 字体与配色（品牌中性默认值，可按需替换）
FONT_NAME = "微软雅黑"
TITLE_COLOR = (31, 56, 100)      # 深蓝
BODY_COLOR = (51, 51, 51)        # 深灰
SUBTITLE_COLOR = (89, 89, 89)    # 浅灰
ACCENT_COLOR = (0, 112, 192)     # 蓝

_PPT_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def build_slide_outline_prompt(report_markdown: str, figures_summary: str) -> str:
    """构建幻灯片大纲提炼的提示词。

    要求 LLM 把段落浓缩成要点、只引用给定清单里的图表文件名（防止编造路径）、
    输出结构化 YAML。
    """
    return f"""你是资深的数据分析演示文稿设计师。请把下面的数据分析报告浓缩成一份用于汇报演示的幻灯片大纲（提炼要点，不要整段照抄原文）。

要求：
1. 只输出 YAML，结构如下（type 只能是 title / section / figure / conclusion 四种）：
```yaml
title: "演示标题"
subtitle: "副标题"
slides:
  - type: "title"
    title: "一页标题"
    subtitle: "副标题"
  - type: "section"
    title: "一、数据概览"
    bullets:
      - "一句话要点"
  - type: "figure"
    title: "营收趋势"
    figure: "确切的图表文件名"
    caption: "图表说明"
  - type: "conclusion"
    title: "总结"
    bullets:
      - "核心结论"
```
2. section 的 bullets 控制在 3~6 条，每条一句话；把段落浓缩成要点，不要照抄。
3. figure 幻灯片的 figure 字段只能使用下面「可用图表」清单里给出的文件名（精确匹配），不要编造路径；没有合适的图表就不要输出 figure 类型。
4. 整体控制在 6~12 页。

📄 报告正文：
{report_markdown}

🖼 可用图表：
{figures_summary}
"""


def parse_slide_outline(response_text: str) -> list[dict[str, Any]]:
    """解析 LLM 返回的幻灯片大纲 YAML，归一化并丢弃非法项。

    复用 protocol.py 的 parse_yaml_response（已处理 LLM 漏闭合围栏的情况）。
    """
    data = parse_yaml_response(response_text)

    slides = data.get("slides") if isinstance(data, dict) else None
    if not isinstance(slides, list) and isinstance(data, list):
        slides = data
    if not isinstance(slides, list):
        return []

    allowed_types = {"title", "section", "figure", "conclusion"}
    valid: list[dict[str, Any]] = []
    for s in slides:
        if not isinstance(s, dict):
            continue
        stype = s.get("type")
        if stype not in allowed_types:
            continue
        title = str(s.get("title") or "").strip()
        if not title and stype != "title":
            continue

        item: dict[str, Any] = {"type": stype, "title": title}
        if stype == "title":
            item["subtitle"] = str(s.get("subtitle") or "").strip()
        elif stype in ("section", "conclusion"):
            bullets = s.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            item["bullets"] = [str(b).strip() for b in bullets if str(b).strip()]
        elif stype == "figure":
            item["figure"] = str(s.get("figure") or "").strip()
            item["caption"] = str(s.get("caption") or "").strip()
            if not item["figure"]:
                continue  # figure 类型必须有文件名，否则无意义
        valid.append(item)

    return valid


def _fallback_outline(report_markdown: str) -> list[dict[str, Any]]:
    """LLM 不可用时的最小确定性大纲：标题页 + 按 ## 标题生成的若干 section。"""
    outline: list[dict[str, Any]] = [{"type": "title", "title": "数据分析报告", "subtitle": ""}]

    sections: list[dict[str, Any]] = []
    for raw in report_markdown.splitlines():
        line = raw.strip()
        if line.startswith("## "):
            sections.append({"type": "section", "title": line[3:].strip(), "bullets": []})
        elif sections and line and not line.startswith("#") and not line.startswith("!"):
            cur = sections[-1]
            if len(cur["bullets"]) < 5:
                cur["bullets"].append(line[:120])

    sections = [s for s in sections if s["bullets"]]
    if sections:
        outline.extend(sections[:8])
    else:
        outline.append({"type": "section", "title": "报告内容", "bullets": ["详见最终分析报告"]})

    return outline


async def build_slide_outline(
    llm,
    report_markdown: str,
    figures: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """调用 LLM 提炼幻灯片大纲；失败时回退到确定性大纲。"""
    figures_summary = build_figures_summary(figures)
    prompt = build_slide_outline_prompt(report_markdown, figures_summary)

    try:
        from langchain_core.messages import HumanMessage, SystemMessage

        response = await llm.ainvoke(
            [
                SystemMessage(content="你是数据分析演示文稿大纲设计助手，只输出 YAML，不要额外解释。"),
                HumanMessage(content=prompt),
            ]
        )
        outline = parse_slide_outline(str(response.content))
        if outline:
            return outline
    except Exception as e:
        print(f"⚠️ 幻灯片大纲提炼失败（回退到确定性大纲）: {e}")

    return _fallback_outline(report_markdown)


def _set_run_font(
    run,
    name: str = FONT_NAME,
    size: Optional[int] = None,
    bold: Optional[bool] = None,
    color: Optional[tuple[int, int, int]] = None,
) -> None:
    """统一设置 run 的字体、字号、加粗、颜色。

    python-pptx 的 ``run.font.name`` 只写 ``a:latin``（拉丁字体），中文需要额外
    写 ``a:ea``（东亚字体）的 typeface，否则中文字符会用默认字体渲染。
    """
    font = run.font
    font.name = name
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = RGBColor(*color)

    try:
        from lxml import etree

        from pptx.oxml.ns import qn

        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(rPr, qn("a:ea"))
        ea.set("typeface", name)
        rPr.set("lang", "zh-CN")
    except Exception:
        pass


def _add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _blank_layout_idx(prs) -> int:
    for i, layout in enumerate(prs.slide_layouts):
        if layout.name and "blank" in layout.name.lower():
            return i
    return 6


def _render_title_slide(prs, layout, spec: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(layout)
    title = spec.get("title") or "数据分析报告"
    subtitle = spec.get("subtitle") or ""

    tf = _add_textbox(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.4))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = title
    _set_run_font(run, size=40, bold=True, color=TITLE_COLOR)

    if subtitle:
        tf2 = _add_textbox(slide, Inches(0.8), Inches(4.2), Inches(11.7), Inches(0.8))
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = subtitle
        _set_run_font(run2, size=20, color=SUBTITLE_COLOR)


def _render_section_slide(prs, layout, spec: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(layout)
    title = spec.get("title") or ""
    bullets = spec.get("bullets") or []

    tf = _add_textbox(slide, Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.9))
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run_font(run, size=32, bold=True, color=TITLE_COLOR)

    # 标题下的强调色短线
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.45), Inches(2.2), Pt(3))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*ACCENT_COLOR)
    bar.line.fill.background()

    body = _add_textbox(slide, Inches(0.95), Inches(1.95), Inches(11.4), Inches(5.0))
    for i, b in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.space_after = Pt(14)
        run = p.add_run()
        text = str(b).strip()
        run.text = text if text.startswith(("•", "-", "·")) else f"• {text}"
        _set_run_font(run, size=20, color=BODY_COLOR)


def _render_figure_slide(prs, layout, spec: dict[str, Any], figures_dir: str) -> None:
    slide = prs.slides.add_slide(layout)
    title = spec.get("title") or ""
    caption = spec.get("caption") or ""
    filename = spec.get("figure") or ""

    tf = _add_textbox(slide, Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run_font(run, size=30, bold=True, color=TITLE_COLOR)

    # 图片路径与 figures 端点一致：仅取 basename 防穿越，图片在会话输出目录下
    img_path = os.path.join(figures_dir, os.path.basename(filename))
    if os.path.isfile(img_path):
        try:
            pic = slide.shapes.add_picture(img_path, Inches(1.0), Inches(1.6), height=Inches(4.4))
            pic.left = int((prs.slide_width - pic.width) / 2)
        except Exception as e:
            print(f"⚠️ 幻灯片插入图片失败（{img_path}）: {e}")

    if caption:
        tf2 = _add_textbox(slide, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.7))
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.CENTER
        run2 = p2.add_run()
        run2.text = caption
        _set_run_font(run2, size=16, color=SUBTITLE_COLOR)


def render_pptx(
    outline: list[dict[str, Any]],
    figures_dir: str,
    output_path: str,
    title: str = "数据分析报告",
) -> str:
    """渲染幻灯片大纲为 .pptx 文件（16:9），返回输出路径。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[_blank_layout_idx(prs)]

    # 确保最前面有一页标题页
    if not outline or outline[0].get("type") != "title":
        outline = [{"type": "title", "title": title, "subtitle": ""}] + list(outline)

    for spec in outline:
        stype = spec.get("type")
        if stype == "title":
            _render_title_slide(prs, layout, spec)
        elif stype in ("section", "conclusion"):
            _render_section_slide(prs, layout, spec)
        elif stype == "figure":
            _render_figure_slide(prs, layout, spec, figures_dir)

    # 兜底：大纲被全部丢弃时也至少产出一页
    if len(prs.slides) == 0:
        _render_title_slide(prs, layout, {"type": "title", "title": title, "subtitle": ""})

    prs.save(output_path)
    return output_path
