"""Word 文档生成逻辑测试。"""

import base64
import os

from docx import Document

from daa.core.docx_gen import render_docx

# 1x1 透明 PNG（Word 插入图片所需的最小合法文件）
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def _write_png(path: str) -> None:
    with open(path, "wb") as f:
        f.write(_PNG_1X1)


def test_render_docx(tmp_path):
    """Markdown 渲染成合法 .docx：标题/段落/内联/列表/图片/表格齐全。"""
    figures_dir = str(tmp_path)
    _write_png(os.path.join(figures_dir, "chart.png"))

    markdown = """# 营收分析报告

## 一、概述

这是**加粗**的段落，包含 `代码` 内容。

- 要点一
- 要点二

1. 第一
2. 第二

![营收趋势](./chart.png)

| 指标 | 数值 |
| --- | --- |
| 营收 | 100 |
"""
    out_path = os.path.join(figures_dir, "最终分析报告.docx")
    result = render_docx(markdown, figures_dir, out_path)

    assert result == out_path
    assert os.path.isfile(out_path)

    doc = Document(out_path)
    body_text = "\n".join(p.text for p in doc.paragraphs)

    # 标题与段落
    assert "营收分析报告" in body_text
    # 内联加粗 / 代码被正确拆出
    assert "加粗" in body_text
    assert "代码" in body_text
    # 列表
    assert "要点一" in body_text
    assert "第一" in body_text
    # 表格
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text == "营收"
    # 图片
    assert len(doc.inline_shapes) == 1


def test_render_docx_missing_figure_does_not_crash(tmp_path):
    """引用了不存在的图片时静默跳过，仍产出合法 .docx。"""
    figures_dir = str(tmp_path)
    markdown = "# 报告\n\n![不存在](./missing.png)\n\n正文内容。\n"
    out_path = os.path.join(figures_dir, "out.docx")
    render_docx(markdown, figures_dir, out_path)

    doc = Document(out_path)
    assert len(doc.inline_shapes) == 0
    assert "正文内容" in "\n".join(p.text for p in doc.paragraphs)
