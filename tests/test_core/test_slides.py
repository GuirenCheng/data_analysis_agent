"""PPT 幻灯片生成逻辑测试。"""

import base64
import os

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from daa.core.slides import parse_slide_outline, render_pptx

# 1x1 透明 PNG（figure 幻灯片插入图片所需的最小合法文件）
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def _write_png(path: str) -> None:
    with open(path, "wb") as f:
        f.write(_PNG_1X1)


def test_parse_slide_outline_with_fence():
    """解析带闭合围栏的大纲 YAML。"""
    response = """```yaml
title: "营业趋势分析"
subtitle: "基于营收数据"
slides:
  - type: "title"
    title: "营业趋势分析"
    subtitle: "基于营收数据"
  - type: "section"
    title: "一、数据概览"
    bullets:
      - "营收逐年增长"
      - "成本同步上升"
  - type: "figure"
    title: "营收趋势"
    figure: "trend.png"
    caption: "近五年营收趋势"
  - type: "conclusion"
    title: "总结"
    bullets: ["建议加大投入"]
```"""
    outline = parse_slide_outline(response)
    assert len(outline) == 4
    assert outline[0]["type"] == "title"
    assert outline[1]["bullets"] == ["营收逐年增长", "成本同步上升"]
    assert outline[2]["figure"] == "trend.png"


def test_parse_slide_outline_without_closing_fence():
    """LLM 漏掉闭合围栏时仍能解析（复用 protocol 的兜底）。"""
    response = """```yaml
title: "报告"
slides:
  - type: "section"
    title: "结论"
    bullets:
      - "要点A"
      - "要点B"
"""
    outline = parse_slide_outline(response)
    assert len(outline) == 1
    assert outline[0]["title"] == "结论"


def test_parse_slide_outline_drops_invalid():
    """缺文件名的 figure 与未知 type 被丢弃，仅保留合法项。"""
    response = """```yaml
slides:
  - type: "figure"
    title: "缺少文件名"
  - type: "unknown_type"
    title: "非法类型"
  - type: "section"
    title: "合法"
    bullets: ["a"]
```"""
    outline = parse_slide_outline(response)
    assert [s["type"] for s in outline] == ["section"]


def test_render_pptx(tmp_path):
    """渲染出合法 .pptx，幻灯片数量正确且 figure 页含图片。"""
    figures_dir = str(tmp_path)
    img_name = "trend.png"
    _write_png(os.path.join(figures_dir, img_name))

    outline = [
        {"type": "title", "title": "数据分析报告", "subtitle": "副标题"},
        {"type": "section", "title": "一、数据概览", "bullets": ["要点1", "要点2"]},
        {"type": "figure", "title": "营收趋势", "figure": img_name, "caption": "趋势"},
        {"type": "conclusion", "title": "总结", "bullets": ["结论"]},
    ]
    out_path = os.path.join(figures_dir, "最终分析报告.pptx")
    result = render_pptx(outline, figures_dir, out_path, title="数据分析报告")

    assert result == out_path
    assert os.path.isfile(out_path)

    prs = Presentation(out_path)
    assert len(prs.slides) == 4

    figure_slide = prs.slides[2]
    pictures = [s for s in figure_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1


def test_render_pptx_missing_figure_does_not_crash(tmp_path):
    """figure 引用了不存在的图片时静默跳过，不崩溃。"""
    figures_dir = str(tmp_path)
    outline = [
        {"type": "title", "title": "报告", "subtitle": ""},
        {"type": "figure", "title": "营收趋势", "figure": "不存在.png", "caption": "x"},
    ]
    out_path = os.path.join(figures_dir, "out.pptx")
    render_pptx(outline, figures_dir, out_path)

    prs = Presentation(out_path)
    assert len(prs.slides) == 2
    pictures = [s for s in prs.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 0
